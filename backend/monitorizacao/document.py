"""
POQ-Impacto (.docx) following the ICNF model "Plano Operacional de Queima — POQ - Impacto": per stratum
(manta morta, herbáceas, arbustos, árvores, restos lenhosos) the 1.ª fase (immediate) and 2.ª fase (6 months)
columns, then the photo index. The permanent annual monitoring (soil, regeneration, parcel-specific criteria)
follows as an annex.
"""
import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from fire_actions.poq import BRAND, DIM, _fmt, _heading, _p, _photo_jpeg, _table
from .models import Assessment


def _yn(v):
    return '' if v is None else 'S' if v else 'N'


def _grouped_table(doc, groups, header, rows, size=8):
    """Table with a first header row of merged group cells, e.g. [('', 1), ('1.ª FASE', 5), ('2.ª FASE', 2)]."""
    t = _table(doc, header, rows, size=size)
    top = t.rows[0]._tr
    new = t.add_row()
    top.addprevious(new._tr)   # move the new row to the top
    col = 0
    for label, span in groups:
        cell = new.cells[col]
        if span > 1:
            cell = cell.merge(new.cells[col + span - 1])
        cell.text = ''
        run = cell.paragraphs[0].add_run(label)
        run.bold = True
        run.font.size = Pt(size)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        col += span
    return t


def generate_poq_impacto(plan):
    action = plan.pre_plan
    executions = [ex for ex in plan.executions.select_related('parcel').prefetch_related(
        'assessments__criterion_values__criterion', 'assessments__photos', 'waivers').order_by('date', 'parcel__name')]
    # Later days of the same parcel are monitored with the first day
    first = {}
    for ex in executions:
        first.setdefault(ex.parcel_id, ex)
    executions = list(first.values())
    numbers = dict(action.pfc.plan_parcels.values_list('parcel_id', 'number')) if action.pfc_id else {}
    responsible = getattr(action.responsible_user, 'operative', None) if action.responsible_user_id else None

    def idp(ex):
        return str(numbers.get(ex.parcel_id, ex.parcel.name))

    def phase(ex, code):
        return next((a for a in ex.assessments.all() if a.phase == code), None) or Assessment()

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Cm(1.5), Cm(1.5), Cm(1.8), Cm(1.8)

    # ── Header ──
    _p(doc, 'Plano Operacional de Queima', bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, color=BRAND)
    _p(doc, 'POQ - Impacto', bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    if action.pfc_id:
        _p(doc, f'PLANO DE FOGO CONTROLADO N.º: {action.pfc.number}', bold=True)
    elif action.pfc_reference:
        _p(doc, f'PLANO DE FOGO CONTROLADO N.º: {action.pfc_reference}', bold=True)
    else:   # queimada without PFC: only the parcel designation
        _p(doc, 'Parcela: ' + ', '.join(p.name for p in action.parcels.all()), bold=True)
    _p(doc, f'Técnico Responsável: {responsible.name if responsible else (action.responsible or "—")}')
    _p(doc, f'Credenciado n.º: {responsible.credential_number if responsible and responsible.credential_number else "—"}')
    d0 = min((ex.date for ex in executions), default=plan.execution_date)
    _p(doc, 'Período de queima: ' + ('1 de janeiro a 31 de maio' if d0.month <= 5 else
                                     '1 de outubro a 31 de dezembro' if d0.month >= 10 else d0.strftime('%d/%m')))
    _p(doc, f'Ano de {d0.year}')
    cover = '(0%, <25%, 26-50%, 51-75%, >75%)'

    # ── Manta morta ──
    _heading(doc, 'MANTA MORTA')
    rows = []
    for ex in executions:
        a1, a2 = phase(ex, Assessment.PHASE_IMEDIATA), phase(ex, Assessment.PHASE_6M)
        rows.append([idp(ex), a1.get_litter_area_display(), a1.litter_red_l, a1.litter_red_f, a1.litter_red_h,
                     _fmt(a1.soil_exposed_pct), _yn(a2.soil_movement), _yn(a2.gullying)])
    _grouped_table(doc, [('', 1), ('1.ª FASE', 5), ('2.ª FASE', 2)],
                   ['ID_PARC', f'Área percorrida {cover}', 'Redução espessura L', 'F', 'H',
                    'Exposição do solo mineral %', 'Erosão: movimento de solo (S/N)', 'Ravinação (S/N)'], rows)

    # ── Herbáceas ──
    _heading(doc, 'HERBÁCEAS')
    rows = []
    for ex in executions:
        a1, a2 = phase(ex, Assessment.PHASE_IMEDIATA), phase(ex, Assessment.PHASE_6M)
        rows.append([idp(ex), a1.get_herb_area_display(), _fmt(a1.herb_resprout_pct), _fmt(a2.herb_resprout_pct),
                     _fmt(a2.herb_height_cm), _yn(a2.herb_germination)])
    _grouped_table(doc, [('', 1), ('1.ª FASE', 2), ('2.ª FASE', 3)],
                   ['ID_PARC', f'Área percorrida {cover}', 'Rebentamento % da área', 'Rebentamento % da área',
                    'Altura média (cm)', 'Germinação (S/N)'], rows)

    # ── Arbustos ──
    _heading(doc, 'ARBUSTOS')
    rows = []
    for ex in executions:
        a1, a2 = phase(ex, Assessment.PHASE_IMEDIATA), phase(ex, Assessment.PHASE_6M)
        rows.append([idp(ex), a1.get_shrub_area_display(), _fmt(a1.shrub_crown_consumed_pct),
                     _yn(a1.shrub_resprout), _fmt(a2.shrub_dead_crowns_pct), _fmt(a2.shrub_resprout_pct),
                     _fmt(a2.shrub_height_cm), _yn(a2.shrub_germination)])
    _grouped_table(doc, [('', 1), ('1.ª FASE', 3), ('2.ª FASE', 4)],
                   ['ID_PARC', f'Área percorrida {cover}', 'Copa consumida %', 'Rebentamento (S/N)',
                    'Copas mortas %', 'Rebentamento % da área', 'Altura média (cm)', 'Germinação (S/N)'], rows)

    # ── Árvores ──
    _heading(doc, 'ÁRVORES')
    rows = []
    for ex in executions:
        a1, a2 = phase(ex, Assessment.PHASE_IMEDIATA), phase(ex, Assessment.PHASE_6M)
        mort = _fmt(a2.tree_mortality_pct, ' %')
        if mort and a2.tree_mortality_dbh_cm is not None:
            mort += f' - DAP ≤ {_fmt(a2.tree_mortality_dbh_cm)} cm'
        rows.append([idp(ex), _fmt(a1.tree_scorch_pct), _fmt(a1.tree_scorch_height_m), _fmt(a1.tree_bole_min_m),
                     _fmt(a1.tree_bole_max_m), _yn(a1.tree_resin), mort, _yn(a2.tree_bark_beetles),
                     _yn(a2.tree_regeneration)])
    _grouped_table(doc, [('', 1), ('1.ª FASE', 5), ('2.ª FASE', 3)],
                   ['ID_PARC', 'Crestamento % da copa', 'Crestamento altura (m)', 'Tronco queimado mín. (m)',
                    'Tronco queimado máx. (m)', 'Derrame de resina (S/N)', 'Mortalidade % em árv. com DAP até',
                    'Por escolitídeos (S/N)', 'Regeneração (S/N)'], rows)

    # ── Restos lenhosos e observações ──
    _heading(doc, 'RESTOS LENHOSOS DE EXPLORAÇÃO e OBSERVAÇÕES GERAIS')
    rows = []
    for ex in executions:
        a1, a2 = phase(ex, Assessment.PHASE_IMEDIATA), phase(ex, Assessment.PHASE_6M)
        obs = '\n'.join(f'{label}: {a.notes}' for label, a in (('1.ª fase', a1), ('2.ª fase', a2)) if a.notes)
        rows.append([idp(ex), a1.get_debris_area_display(), obs])
    _grouped_table(doc, [('', 1), ('1.ª FASE', 1), ('2.ª FASE', 1)],
                   ['ID_PARC', f'Área percorrida {cover}', 'Observações gerais'], rows)

    # ── Índice de fotografia ──
    _heading(doc, 'Índice de fotografia')
    photos = [(ex, a, ph) for ex in executions for a in sorted(ex.assessments.all(), key=lambda x: x.date)
              for ph in a.photos.all()]
    rows = [[idp(ex), a.date.strftime('%d/%m/%Y'), f'{a.label}' + (f' — {ph.caption}' if ph.caption else ''),
             f'Foto {i}'] for i, (ex, a, ph) in enumerate(photos, 1)]
    _table(doc, ['ID_PARC', 'DATA', 'Descrição', 'Identificação da FOTO'], rows)

    # ── Anexo: monitorização permanente ──
    visits = [(ex, a) for ex in executions for a in sorted(ex.assessments.all(), key=lambda x: x.date)]
    if visits:
        _heading(doc, 'ANEXO — Monitorização (solo e regeneração da vegetação)')
        rows = [[idp(ex), a.label, a.date.strftime('%d/%m/%Y'), _fmt(a.soil_cover_pct), _fmt(a.soil_exposed_pct),
                 _yn(a.soil_movement), _yn(a.gullying), _fmt(a.soil_ph), _fmt(a.soil_organic_matter_pct),
                 _fmt(a.veg_cover_pct), a.veg_species] for ex, a in visits]
        _table(doc, ['ID_PARC', 'Visita', 'Data', 'Cobertura do solo %', 'Solo mineral exposto %',
                     'Movimento de solo', 'Ravinação', 'pH', 'Matéria orgânica %', 'Cobertura vegetal %',
                     'Espécies em regeneração'], rows, size=7.5)
        crit = [[idp(ex), a.label, a.date.strftime('%d/%m/%Y'), str(v.criterion), v.value]
                for ex, a in visits for v in a.criterion_values.all()]
        if crit:
            _p(doc, 'Critérios específicos das parcelas', bold=True, size=9).paragraph_format.space_before = Pt(8)
            _table(doc, ['ID_PARC', 'Visita', 'Data', 'Critério', 'Valor'], crit)
        notes = [(idp(ex), a) for ex, a in visits if a.soil_notes]
        for pid, a in notes:
            _p(doc, f'{pid} — {a.label}: {a.soil_notes}', size=8.5, color=DIM)

    waivers = [(ex, w) for ex in executions for w in ex.waivers.all()]
    if waivers:
        _heading(doc, 'Visitas dispensadas')
        _table(doc, ['ID_PARC', 'Visita', 'Dispensada em', 'Motivo'],
               [[idp(ex), w.label, w.created_at.strftime('%d/%m/%Y'), w.reason] for ex, w in waivers])

    # ── Fotografias ──
    if photos:
        _heading(doc, 'Fotografias')
        for i, (ex, a, ph) in enumerate(photos, 1):
            buf = _photo_jpeg(ph)
            if buf:
                doc.add_picture(buf, width=Cm(12))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                _p(doc, f'Foto {i} — {idp(ex)} · {a.label}' + (f' — {ph.caption}' if ph.caption else ''),
                   size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=DIM)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
