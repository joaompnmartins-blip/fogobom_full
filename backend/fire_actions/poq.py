"""
POQ-Execução (.docx) following the ICNF model "Plano Operacional de Queima — POQ - Execução":
technicians involved, operational problems, fuel moisture, weather, conduction scheme per parcel,
immediate effects, efficacy, economic evaluation, photo index, incident (art. 15.º), and the annexes
(difficulty classification, field sheets for matos / povoamentos).
"""
import io
import json
import math
from collections import OrderedDict

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont, ImageOps

from .models import CONDUCT, RATINGS, WEATHER_CODES

BRAND = RGBColor(0xD0, 0x5C, 0x1E)
DIM = RGBColor(0x66, 0x66, 0x66)


def _p(doc, text='', bold=False, size=10, align=None, color=None, italic=False, space_after=3):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    r.font.name = 'Arial'
    if color:
        r.font.color.rgb = color
    return p


def _heading(doc, text):
    p = _p(doc, text, bold=True, size=12, color=BRAND, space_after=5)
    p.paragraph_format.space_before = Pt(12)
    return p


def _table(doc, header, rows, size=8.5):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ''
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(size)
    for row in rows or [['—'] * len(header)]:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            for k, line in enumerate(str('' if v is None else v).split('\n')):
                para = cells[i].paragraphs[0] if k == 0 else cells[i].add_paragraph()
                run = para.add_run(line)
                run.font.size = Pt(size)
    return t


def _fmt(v, unit=''):
    if v is None or v == '':
        return ''
    if isinstance(v, float):
        v = f'{v:g}'
    return f'{v}{unit}'


def _id_parc(ex, numbers):
    return str(numbers.get(ex.parcel_id, ex.parcel.name))


def _parcel_schema_png(parcel, width=900, height=520, pad=50):
    """Clean drawing of the parcel boundary (used when no conduction scheme was uploaded)."""
    if not parcel.geometry:
        return None
    ring = json.loads(parcel.geometry.json)['coordinates'][0]
    lngs, lats = [c[0] for c in ring], [c[1] for c in ring]
    lat0 = sum(lats) / len(lats)
    kx = math.cos(math.radians(lat0))
    xs, ys = [x * kx for x in lngs], lats
    xr, yr = (max(xs) - min(xs)) or 1e-6, (max(ys) - min(ys)) or 1e-6
    scale = min((width - 2 * pad) / xr, (height - 2 * pad) / yr)
    ox = pad + ((width - 2 * pad) - xr * scale) / 2
    oy = pad + ((height - 2 * pad) - yr * scale) / 2
    pts = [(ox + (x - min(xs)) * scale, height - (oy + (y - min(ys)) * scale)) for x, y in zip(xs, ys)]
    img = Image.new('RGB', (width, height), 'white')
    d = ImageDraw.Draw(img)
    d.polygon(pts, fill='#fff3ec', outline=None)
    d.line(pts + [pts[0]], fill='#D05C1E', width=3)
    try:
        font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', 16)
    except Exception:
        font = ImageFont.load_default()
    d.text((width - 40, 15), 'N ↑', fill='#333', font=font)
    d.text((pad, 15), parcel.name, fill='#D05C1E', font=font)
    m_per_px = 111320 / scale
    bar_m = 10 ** math.floor(math.log10(max(m_per_px * 150, 1)))
    bar_px = bar_m / m_per_px
    d.line([(pad, height - 25), (pad + bar_px, height - 25)], fill='#444', width=2)
    d.text((pad, height - 22), f'{bar_m:g} m', fill='#444', font=font)
    buf = io.BytesIO()
    img.save(buf, 'PNG')
    buf.seek(0)
    return buf


def _photo_jpeg(photo, max_px=1400):
    try:
        with photo.image.open('rb') as f:
            img = Image.open(f)
            img.load()
    except Exception:
        return None
    img = ImageOps.exif_transpose(img).convert('RGB')
    img.thumbnail((max_px, max_px))
    buf = io.BytesIO()
    img.save(buf, 'JPEG', quality=85)
    buf.seek(0)
    return buf


def generate_poq_execucao(plan):
    action = plan.pre_plan
    executions = list(plan.executions.select_related('parcel').prefetch_related(
        'members__operative', 'readings', 'strips', 'cost_lines').order_by('date', 'parcel__name'))
    numbers = {}
    if action.pfc_id:
        numbers = dict(action.pfc.plan_parcels.values_list('parcel_id', 'number'))
    responsible = getattr(action.responsible_user, 'operative', None) if action.responsible_user_id else None

    doc = Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(1.8), Cm(1.8), Cm(2), Cm(2)

    # ── Header ──
    _p(doc, 'Plano Operacional de Queima', bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, color=BRAND)
    _p(doc, 'POQ - Execução', bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    if action.pfc_id:
        _p(doc, f'PLANO DE FOGO CONTROLADO N.º: {action.pfc.number}', bold=True)
    elif action.pfc_reference:
        _p(doc, f'PLANO DE FOGO CONTROLADO N.º: {action.pfc_reference}', bold=True)
    else:   # queimada without PFC: only the parcel designation
        _p(doc, 'Parcela: ' + ', '.join(p.name for p in action.parcels.all()), bold=True)
    _p(doc, f'Técnico Responsável de Queima: {responsible.name if responsible else (action.responsible or "—")}')
    _p(doc, f'Credenciado n.º: {responsible.credential_number if responsible and responsible.credential_number else "—"}')
    dates = sorted({ex.date for ex in executions}) or [plan.execution_date]
    d0 = dates[0]
    period = (f'De 1 de Janeiro a 31 de Maio de {d0.year}' if d0.month <= 5
              else f'De 1 de Outubro a 31 de Dezembro de {d0.year}' if d0.month >= 10
              else f'Fora dos períodos de queima habituais ({d0:%d/%m/%Y})')
    _p(doc, f'Período de queima: {period}')
    _p(doc, f'{action.get_action_type_display()}'
            + (f' — {action.get_purpose_display()}' if action.purpose else '')
            + (' · ação de formação' if action.training else ''), size=9, color=DIM)

    # ── Técnicos envolvidos ──
    _heading(doc, 'Técnicos envolvidos na queima')
    rows = []
    for ex in executions:
        for m in ex.members.all():
            rows.append([_id_parc(ex, numbers), f'{ex.date:%d/%m/%Y}', m.operative.name,
                         m.operative.credential_number or '—', m.get_function_display() or '—'])
    _table(doc, ['ID_PARC', 'DATA', 'NOME DO TÉCNICO', 'N.º DE CREDENCIAÇÃO', 'FUNÇÕES'], rows)

    # ── Problemas ──
    _heading(doc, 'Problemas operacionais')
    rows = [[_id_parc(ex, numbers), f'{ex.date:%d/%m/%Y}',
             '\n'.join(filter(None, [ex.problems_display, ex.problems_description])) or 'Sem problemas']
            for ex in executions]
    _table(doc, ['ID_PARC', 'DATA', 'DESCRIÇÃO'], rows)

    # ── Humidade do combustível ──
    _heading(doc, 'Humidade do combustível')
    rows = [[_id_parc(ex, numbers), f'{ex.date:%d/%m/%Y}', ex.fuel_method or '', ex.fuel_superficial,
             ex.fuel_manta_f, ex.fuel_manta_h] for ex in executions]
    _table(doc, ['ID_PARC', 'DATA', 'Metodologia', 'Superficial', 'Manta morta F', 'Manta morta H'], rows)
    _p(doc, 'Metodologia: A – amostra, B – tabela do guia, C – BEHAVE, D – índice do Sistema Canadiano. '
            'Valores em % ou classe (A muito húmido … E bastante seco) ou FFMC/DMC.', size=8, color=DIM)

    # ── Meteorologia ──
    _heading(doc, 'Meteorologia')
    rows = []
    for ex in executions:
        for r in ex.readings.all():
            wind = ' / '.join(filter(None, [_fmt(r.wind_bft, ' Bft'), _fmt(r.wind_kmh, ' km/h')]))
            rows.append([_id_parc(ex, numbers), f'{ex.date:%d/%m/%Y}', r.time.strftime('%H:%M'), _fmt(r.temp),
                         _fmt(r.rh), _fmt(r.weather), wind, r.wind_dir,
                         _fmt(r.conduct) + (f' ({r.conduct_other})' if r.conduct_other else '')])
    _table(doc, ['ID_PARC', 'Data', 'Hora', 'Temp. ar, °C', 'Hum. relat., %', 'Est. do tempo *', 'Vel. vento **',
                 'Dir. vento ***', 'Conduç. ****'], rows, size=8)
    _p(doc, 'Dados no início e fim da queima e sempre que a alteração de algum fator meteorológico provoque '
            'alterações consideráveis no comportamento.', size=8, color=DIM)
    _p(doc, '* Estado do tempo: ' + '; '.join(l for _, l in WEATHER_CODES) + '.', size=8, color=DIM)
    _p(doc, '** Velocidade do vento a 2 m (Beaufort, km/h). *** Direção do vento (N, NE, E, SE, S, SW, W, NW, V = variável).',
       size=8, color=DIM)
    _p(doc, '**** Condução do fogo: ' + '; '.join(l for _, l in CONDUCT) + '.', size=8, color=DIM)

    # ── Esquema de condução ──
    _heading(doc, 'Esquema de condução da queima')
    for ex in executions:
        _table(doc, ['ID_PARC', 'Data', 'Início', 'Fim', 'Téc. n.º', 'VLCI (< 500 l)', 'Homens', 'VFCI'],
               [[_id_parc(ex, numbers), f'{ex.date:%d/%m/%Y}', ex.start_time.strftime('%H:%M') if ex.start_time else '',
                 ex.end_time.strftime('%H:%M') if ex.end_time else '', _fmt(ex.n_tecnicos), _fmt(ex.n_vlci),
                 _fmt(ex.n_men), _fmt(ex.n_vfci) + (f' ({ex.vfci_capacity_l} l)' if ex.vfci_capacity_l else '')]])
        img = None
        if ex.schema:
            try:
                with ex.schema.open('rb') as f:
                    img = io.BytesIO(f.read())
            except Exception:
                img = None
        img = img or _parcel_schema_png(ex.parcel)
        if img:
            doc.add_picture(img, width=Cm(15))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if ex.operation_notes:
            _p(doc, ex.operation_notes, size=9)
        _p(doc, 'ESQUEMA: sentidos do declive, Norte, direção do fogo, localização da ignição (A), das faixas de '
                'segurança (B) e da fotografia (C).', size=8, color=DIM)
    _p(doc, 'VLCI – Veículo Ligeiro de Combate a Incêndio; VFCI – Veículo Florestal de Combate a Incêndio.',
       size=8, color=DIM)

    # ── Efeitos imediatos ──
    _heading(doc, 'Efeitos imediatos da queima')
    rows = [[_id_parc(ex, numbers), _fmt(ex.eff_veg_area_pct), _fmt(ex.eff_crowns_pct), _fmt(ex.eff_diameter_mm),
             _fmt(ex.eff_litter_area_pct), _fmt(ex.eff_litter_reduction_pct)] for ex in executions]
    _table(doc, ['ID_PARC', 'Vegetação: área percorrida %', 'Copas consumidas %', 'Diâmetro calcinado mm',
                 'Manta morta: área percorrida %', 'Redução espessura %'], rows)

    # ── Eficácia ──
    _heading(doc, 'Eficácia da queima')
    rating = dict(RATINGS)
    rows = [[_id_parc(ex, numbers), rating.get(ex.rating_objectives, ''), rating.get(ex.rating_fuel, ''),
             rating.get(ex.rating_weather, ''), ex.rating_why] for ex in executions]
    _table(doc, ['ID_PARC', 'Resposta aos objetivos', 'Redução de combustível', 'Condições meteorológicas',
                 'Porquê?'], rows)

    # ── Avaliação económica ──
    _heading(doc, 'Avaliação económica')
    rows = []
    for ex in executions:
        for c in ex.cost_lines.all():
            rows.append([_id_parc(ex, numbers), c.get_category_display(), c.description, _fmt(c.qty_burn),
                         _fmt(c.qty_extinction), _fmt(c.quantity), _fmt(c.unit_price), _fmt(c.amount, ' €')])
        if ex.cost_lines.exists():
            rows.append([_id_parc(ex, numbers), 'TOTAL', '', '', '', '', '', _fmt(ex.cost_total, ' €')])
    _table(doc, ['ID_PARC', 'Categoria', 'Tipo', 'Queima', 'Extinção, vigilância', 'Total (h / l / km)',
                 'Custo unitário €', 'Total'], rows, size=8)
    _p(doc, 'Custos em pessoal: n.º homens × horas, tempo total dedicado à queima (transporte, tempos mortos e '
            'refeições). Custos em material: n.º máquinas × horas, combustível (l) e viaturas (km).', size=8, color=DIM)

    # ── Índice de fotografia ──
    _heading(doc, 'Índice de fotografia')
    photos = list(plan.photos.select_related('parcel').order_by('parcel__name', 'taken_on', 'moment', 'pk'))
    rows = [[numbers.get(ph.parcel_id, ph.parcel.name if ph.parcel else '—'),
             ph.taken_on.strftime('%d/%m/%Y') if ph.taken_on else '', ph.get_moment_display() or '',
             f'Foto {i}' + (f' — {ph.caption}' if ph.caption else '')] for i, ph in enumerate(photos, 1)]
    _table(doc, ['ID_PARC', 'DATA', 'MOMENTO', 'Identificação da FOTO'], rows)

    # ── Incidente (art. 15.º) ──
    if plan.incident:
        _heading(doc, 'Fuga do fogo / incêndio (art. 15.º do Regulamento do Fogo Técnico)')
        _table(doc, ['Campo', 'Descrição'], [
            ['Data', plan.incident_date.strftime('%d/%m/%Y') if plan.incident_date else ''],
            ['Circunstâncias', plan.incident_circumstances], ['Entidades contactadas', plan.incident_entities],
            ['Equipas complementares', plan.incident_teams], ['Responsável pelo combate', plan.incident_commander],
            ['Estratégia de controlo', plan.incident_strategy], ['Área afetada (ha)', _fmt(plan.incident_area_ha)],
            ['Possíveis prejuízos', plan.incident_damages]])
    if plan.notes:
        _heading(doc, 'Observações')
        _p(doc, plan.notes)

    # ── Fotografias ──
    if photos:
        _heading(doc, 'Fotografias')
        for i, ph in enumerate(photos, 1):
            buf = _photo_jpeg(ph)
            if buf:
                doc.add_picture(buf, width=Cm(12))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                _p(doc, f'Foto {i}' + (f' — {ph.caption}' if ph.caption else ''), size=8.5,
                   align=WD_ALIGN_PARAGRAPH.CENTER, color=DIM)

    # ── Anexo: dificuldade ──
    _heading(doc, 'ANEXO — Classificação da dificuldade da operação')
    _p(doc, '(Ferramenta a utilizar pelo técnico, mas não devolver preenchido)', size=8.5, italic=True, color=DIM)
    for parcel in OrderedDict((ex.parcel_id, ex.parcel) for ex in executions).values():
        ev = parcel.evaluations.filter(completed_at__isnull=False).order_by('-completed_at').first()
        if not ev:
            continue
        rows = [[f'{label} ({w})', _fmt(getattr(ev, f)), _fmt(getattr(ev, f) * w if getattr(ev, f) else None)]
                for f, label, w in ev.DIFFICULTY_ELEMENTS]
        rows.append(['TOTAL', '', f'{ev.difficulty_total or "—"} — {ev.difficulty_class or "incompleta"}'])
        _p(doc, f'{numbers.get(parcel.pk, parcel.name)} — {parcel.name}', bold=True, size=9)
        _table(doc, ['Elemento (fator de ponderação)', 'Valor atribuído', 'Peso'], rows)
    _p(doc, 'Fácil / simples < 35 · Normal 35–54 · Difícil / complexa ≥ 55.', size=8, color=DIM)

    # ── Anexo: fichas de campo ──
    sec = doc.add_section()
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = sec.right_margin = Cm(1.5)
    for ex in executions:
        _field_sheet(doc, ex, numbers, action)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _field_sheet(doc, ex, numbers, action):
    parcel = ex.parcel
    kind = 'MATOS' if ex.formation == 'mato' else 'POVOAMENTOS' if ex.formation == 'pov' else 'RESÍDUOS'
    _p(doc, f'ANEXO – FICHA DE CAMPO – {kind} (não é para enviar)', bold=True, size=11, color=BRAND)
    ev = parcel.evaluations.filter(completed_at__isnull=False).order_by('-completed_at').first()
    pfc = action.pfc.number if action.pfc_id else (action.pfc_reference or '—')
    members = list(ex.members.all())
    lead = next((m for m in members if m.function == 'coordenacao'), members[0] if members else None)
    _table(doc, ['PFC n.º', 'Parcela n.º', 'Data', 'Técnico responsável', 'Creden. n.º', 'Técnicos que apoiaram'],
           [[pfc, _id_parc(ex, numbers), f'{ex.date:%d/%m/%Y}', lead.operative.name if lead else '',
             lead.operative.credential_number if lead else '',
             ', '.join(f'{m.operative.name} ({m.operative.credential_number})' for m in members
                       if m is not lead and m.operative.credential_number)]], size=8)
    _p(doc, 'Objetivo(s): ' + (action.objectives_display or (ev.objectives_display if ev else '') or '—'), size=9)
    _p(doc, 'CARACTERIZAÇÃO FÍSICA', bold=True, size=9)
    _table(doc, ['Área (ha)', 'Exposição', 'Declive %', 'Altitude m', 'Rochas %', 'Terraços', 'Topografia',
                 'Espessura manta morta (cm)'],
           [[_fmt(ex.area_treated_ha or parcel.area_ha), ev.get_aspect_display() if ev else '',
             _fmt(ev.slope_pct) if ev else '', _fmt(ev.altitude_m) if ev else '', _fmt(ev.rocks_pct) if ev else '',
             ('sim' if ev.terraces else 'não') if ev and ev.terraces is not None else '',
             ev.get_topography_display() if ev else '', _fmt(ex.litter_thickness_cm)]], size=8)
    if ev and ev.vegetation.exists():
        _p(doc, 'VEGETAÇÃO', bold=True, size=9)
        _table(doc, ['Espécie', '% coberto', 'Altura (m)', 'Idade (anos)'],
               [[v.species, _fmt(v.cover_pct), _fmt(v.height_m), _fmt(v.age_years)] for v in ev.vegetation.all()],
               size=8)
    _p(doc, 'PREPARAÇÃO DE FAIXAS DE SEGURANÇA', bold=True, size=9)
    _table(doc, ['Lado', 'Meio (código)', 'Largura (m)', 'Tempo de execução'],
           [[s.get_side_display(), s.get_means_display() if s.means else '', _fmt(s.width_m), s.time_spent]
            for s in ex.strips.all()], size=8)
    _p(doc, 'METEOROLOGIA E COMPORTAMENTO DO FOGO', bold=True, size=9)
    _table(doc, ['Hora', 'Temp. ar', 'HR', 'Vento km/h', 'Vento Bft', 'Direção', 'Comprim. chama (m)',
                 'Ângulo chama', 'Propagação (m/min)'],
           [[r.time.strftime('%H:%M'), _fmt(r.temp), _fmt(r.rh), _fmt(r.wind_kmh), _fmt(r.wind_bft), r.wind_dir,
             _fmt(r.flame_length_m), _fmt(r.flame_angle), _fmt(r.spread_m_min)] for r in ex.readings.all()], size=8)
    _table(doc, ['Dias s/ chuva', 'FWI', 'ISI', 'FFMC', 'BUI', 'DMC', 'DC', 'Humid. comb. superficial'],
           [[_fmt(ex.dry_days), _fmt(ex.fwi), _fmt(ex.isi), _fmt(ex.ffmc), _fmt(ex.bui), _fmt(ex.dmc), _fmt(ex.dc),
             ex.fuel_superficial]], size=8)
    _p(doc, 'Execução: início ' + (ex.start_time.strftime('%H:%M') if ex.start_time else '—') + ' · fim '
       + (ex.end_time.strftime('%H:%M') if ex.end_time else '—')
       + f' · Técnica(s) de ignição: {ex.ignition_techniques or "—"}', size=9)
    _p(doc, f'Problemas e dificuldades: {ex.problems_display or "—"}', size=9)
    _p(doc, f'Efeitos imediatos: área percorrida {_fmt(ex.eff_veg_area_pct, "%") or "—"}, copa consumida '
            f'{_fmt(ex.eff_crowns_pct, "%") or "—"}, diâmetro calcinado {_fmt(ex.eff_diameter_mm, " mm") or "—"}; '
            f'manta morta: área percorrida {_fmt(ex.eff_litter_area_pct, "%") or "—"}, redução '
            f'{_fmt(ex.eff_litter_reduction_pct, "%") or "—"}', size=9)
    _p(doc, f'Eficácia: {ex.rating_why or "—"}', size=9)
    if ex.cost_lines.exists():
        _p(doc, f'Custos: {_fmt(ex.cost_total, " €")}', size=9)
    _p(doc, '')
