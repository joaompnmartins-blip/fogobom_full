"""
Word document of the PFC following the ICNF model (Modelo-Plano-FogoControlado):
cover, location, mandatory contents (art. 12.º), distribution by burning season,
signature, list of annexes and Anexo 3 — quadro de identificação das parcelas.
"""
import io
from collections import defaultdict

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

BRAND = RGBColor(0xD0, 0x5C, 0x1E)
DIM = RGBColor(0x66, 0x66, 0x66)


def _p(doc, text='', bold=False, size=10.5, align=None, color=None, italic=False, space_after=4):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    r.font.name = 'Arial'
    if color:
        r.font.color.rgb = color
    return p


def _heading(doc, text):
    p = _p(doc, text, bold=True, size=12.5, color=BRAND, space_after=6)
    p.paragraph_format.space_before = Pt(12)
    return p


def _text_or_placeholder(doc, text):
    if text and text.strip():
        for para in text.strip().split('\n'):
            _p(doc, para)
    else:
        _p(doc, '(por preencher)', italic=True, color=DIM)


def _table(doc, header, rows, widths=None, size=8.5):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            for k, line in enumerate(str(v).split('\n')):
                para = cells[i].paragraphs[0] if k == 0 else cells[i].add_paragraph()
                run = para.add_run(line)
                run.font.size = Pt(size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def generate_pfc_docx(pfc):
    doc = Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(2), Cm(2), Cm(2.5), Cm(2.5)
    op = getattr(pfc.technician, 'operative', None)
    tech_name = op.name if op and op.name else pfc.technician.get_full_name() or pfc.technician.get_username()
    loc = pfc.location
    plan_parcels = list(pfc.plan_parcels.select_related('parcel'))

    # ── Cover ──
    _p(doc, 'PLANO DE FOGO CONTROLADO', bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, color=BRAND,
       space_after=10)
    _p(doc, f'PFC {pfc.number}', bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, '(n.º sequencial / n.º de credenciação do técnico / ano de aprovação)', size=8.5,
       align=WD_ALIGN_PARAGRAPH.CENTER, color=DIM)
    _p(doc, pfc.name, bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    _p(doc, f'Período de validade: {pfc.validity_start:%d/%m/%Y} – {pfc.validity_end:%d/%m/%Y}',
       align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, f'Técnico(a): {tech_name} (credenciação n.º {pfc.credential_number})', align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, f'Proponente: {pfc.proponent}', align=WD_ALIGN_PARAGRAPH.CENTER)
    if pfc.management_unit:
        _p(doc, pfc.management_unit, align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, pfc.get_dcnf_display(), align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, 'Distrito: ' + (', '.join(loc['distritos']) or '—'), align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, 'Concelho: ' + (', '.join(loc['concelhos']) or '—'), align=WD_ALIGN_PARAGRAPH.CENTER)

    # ── 1. Location ──
    _heading(doc, '1. Localização')
    _table(doc, ['Concelho(s)', 'Freguesia(s)'], [[', '.join(loc['concelhos']) or '—',
                                                  ', '.join(loc['freguesias']) or '—']], widths=[5, 11])
    _p(doc, f'Localização na(s) folha(s) da Carta Militar de Portugal (escala 1:25 000) n.º '
            f'{pfc.military_map_sheets or "…"} (Anexo 1 — Mapa de localização).')
    _p(doc, 'Cartografia no sistema de referência PT-TM06/ETRS89.', size=9, color=DIM)
    _p(doc, f'Período de intervenção: {pfc.validity_start:%d/%m/%Y} – {pfc.validity_end:%d/%m/%Y} (no máximo 5 anos).')
    _p(doc, f'Área total a intervencionar: {pfc.total_area_ha:g} hectares.')

    # ── 2–4. Characterisation ──
    _heading(doc, '2. Caracterização da área de intervenção')
    _p(doc, 'Orografia, uso do solo, formações vegetais e combustíveis', bold=True)
    _text_or_placeholder(doc, pfc.characterization)
    _p(doc, 'História do fogo', bold=True)
    _text_or_placeholder(doc, pfc.fire_history)
    _p(doc, 'Anexo 2 – Mapa das áreas ardidas (com implantação das parcelas)', size=9, color=DIM)
    _p(doc, 'Localização das infraestruturas de DFCI', bold=True)
    _text_or_placeholder(doc, pfc.dfci_infrastructure)

    _heading(doc, '3. Identificação das parcelas individuais a intervir')
    by_season = defaultdict(float)
    for pp in plan_parcels:
        by_season[pp.season or 'sem época definida'] += pp.parcel.area_ha or 0
    _p(doc, 'Distribuição das áreas a intervencionar por épocas de queima (outubro – abril/maio):')
    _table(doc, ['Época de queima', 'Área a intervencionar (ha)'],
           [[s, f'{a:.2f}'] for s, a in sorted(by_season.items())] or [['—', '0']], widths=[6, 6])
    _p(doc, 'Anexo 3 – Quadro de identificação das parcelas individuais. '
            'Anexo 4 – Mapa de localização das parcelas à escala 1:10 000.', size=9, color=DIM)

    _heading(doc, '4. Objetivos específicos do uso do fogo controlado')
    _text_or_placeholder(doc, pfc.objectives)
    _heading(doc, '5. Meios humanos e materiais necessários à execução do PFC')
    _text_or_placeholder(doc, pfc.means)
    _heading(doc, '6. Plano de emergência e de contingência')
    if pfc.emergency_plan.strip():
        _text_or_placeholder(doc, pfc.emergency_plan)
    else:
        _p(doc, 'Opcional para parcelas ou conjunto de parcelas inferiores a 50 ha.' if not pfc.emergency_plan_required
           else '(por preencher — obrigatório: área total ≥ 50 ha)', italic=True, color=DIM)
    _heading(doc, '7. Outros elementos considerados pertinentes ao PFC')
    _text_or_placeholder(doc, pfc.other_elements)

    _p(doc, '')
    _p(doc, 'O(A) Técnico(a)', align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, '_______________________________', align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, tech_name, align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, f'Credenciação n.º {pfc.credential_number}', align=WD_ALIGN_PARAGRAPH.CENTER)

    _heading(doc, 'Anexos')
    for a in ['Anexo 1 – Mapa de localização (implantado na Carta Militar de Portugal à escala 1:25 000)',
              'Anexo 2 – Mapa das áreas ardidas (com implantação das parcelas)',
              'Anexo 3 – Quadro de identificação das parcelas individuais',
              'Anexo 4 – Mapa de localização das parcelas à escala 1:10 000']:
        _p(doc, a)

    # ── Anexo 3 (landscape) ──
    sec = doc.add_section()
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = sec.right_margin = Cm(1.5)
    _p(doc, 'Anexo 3 – Quadro de identificação das parcelas individuais', bold=True, size=12, color=BRAND)
    rows = []
    for pp in plan_parcels:
        presc = '\n'.join(f'{label}: {value}{(" " + unit) if unit else ""}' for label, value, unit in pp.prescription_rows)
        means = '\n'.join(f'{label}: {v}' for label, v in [
            ('N.º de técnicos FC', pp.n_tecnicos), ('N.º de operacionais de queima', pp.n_operacionais),
            ('N.º de equipas de apoio', pp.n_equipas), ('N.º de elementos de apoio', pp.n_elementos),
            ('N.º de VFCI', pp.n_vfci)] if v is not None)
        rows.append([pp.number, pp.season or '—', pp.occupation or pp.parcel.get_vegetation_type_display(),
                     f'{pp.parcel.area_ha or 0:.2f}', pp.objective or '—', presc or '—',
                     pp.get_ignition_display() or '—', means or '—'])
    _table(doc, ['N.º PARC', 'Época de queima', 'Ocupação', 'Área (ha)', 'Objetivo', 'Prescrição',
                 'Técnica de ignição', 'Meios necessários'],
           rows or [['—'] * 8], widths=[1.4, 2.2, 5, 1.6, 2.6, 6, 2.6, 4.2], size=8)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
