import io
import warnings
from datetime import date, time, timedelta
from unittest import mock

from django.contrib.auth.models import User
from django.core.cache import cache
from django.test import TestCase, override_settings
from django.utils import timezone

from core.models import UserPreferences
from fire_actions.models import BurningPlan, ExecutionMember, FireAction, ParcelExecution
from operatives.models import Operative
from parcelas.models import FireParcel
from .alerts import send_deadline_alerts
from .deadlines import all_deadlines, icnf_window_after
from .fieldbook import entries, hours_summary
from .models import AlertLog

SQUARE = 'POLYGON((-8 37.1,-7.999 37.1,-7.999 37.101,-8 37.101,-8 37.1))'


def make_tecnico(username='tec', number='9-2020', valid_until=None):
    user = User.objects.create_user(username, email=f'{username}@example.com', password='x')
    op = Operative.objects.create(user=user, name=f'Técnico {username}', email=user.email, phone='',
                                  certification_level='', credential_type=Operative.CRED_TECNICO,
                                  credential_number=number, credential_status=Operative.STATUS_VALIDADA,
                                  credential_lifetime=valid_until is None, credential_valid_until=valid_until,
                                  field_book_number='L-12')
    return user, op


class Base(TestCase):
    def setUp(self):
        cache.clear()
        self.user, self.op = make_tecnico()
        self.worker = Operative.objects.create(name='Operacional O', email='o@example.com', phone='',
                                               certification_level='', credential_type=Operative.CRED_OPERACIONAL)
        self.parcel = FireParcel.objects.create(name='Cerro', geometry=SQUARE, status='apta', area_ha=2.5,
                                                freguesia='Alferce', concelho='Monchique', distrito='Faro',
                                                vegetation_type='matos')
        self.action = FireAction.objects.create(name='Queimada do Cerro', scheduled_date=date(2026, 1, 20),
                                                responsible_user=self.user, created_by=self.user,
                                                objectives='prevencao')
        self.action.parcels.set([self.parcel])
        self.plan = BurningPlan.objects.create(pre_plan=self.action, execution_date=date(2026, 1, 20))
        self.ex = ParcelExecution.objects.create(plan=self.plan, parcel=self.parcel, date=date(2026, 1, 20),
                                                 start_time=time(9, 0), end_time=time(13, 30), area_treated_ha=2.5,
                                                 n_vfci=1, validator_role='proprietario', validator_name='J. Silva')
        self.m_tec = ExecutionMember.objects.create(execution=self.ex, operative=self.op, function='coordenacao',
                                                    book_page='7')
        self.m_op = ExecutionMember.objects.create(execution=self.ex, operative=self.worker, function='ignicao',
                                                   hours=5)


class FieldbookTests(Base):
    def test_entries_and_hours(self):
        [e] = entries(self.op)
        self.assertEqual((e['book'], e['page'], e['freguesia'], e['distrito'], e['type'], e['function']),
                         ('L-12', '7', 'Alferce', 'Faro', 'Queimada Autor.', 'Responsável de queima'))
        self.assertEqual((e['hours'], e['validated'], e['validation_role'], e['formation']),
                         (4.5, True, 'Proprietário', 'Mato'))
        self.assertIn('Operacional O', e['notes'])
        self.assertIn('1 VFCI', e['notes'])
        s = hours_summary(self.op, today=date(2026, 6, 1))
        self.assertEqual((s['validated'], s['as_responsible'], s['required']), (4.5, 4.5, 150))
        # The operacional is validated by the técnico, not by the owner
        [w] = entries(self.worker)
        self.assertEqual((w['hours'], w['validated']), (5, False))
        self.m_op.validated_by = self.user
        self.m_op.save()
        [w] = entries(self.worker)
        self.assertEqual((w['validated'], w['validation_role'], w['validation_name']),
                         (True, 'Técnico credenciado', 'Técnico tec'))
        self.assertEqual(hours_summary(self.worker)['required'], 100)

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_pages_and_excel(self):
        import openpyxl
        self.client.force_login(self.user)
        r = self.client.get(f'/livro-de-campo/{self.op.pk}/')
        self.assertContains(r, 'Queimada do Cerro')
        self.assertContains(r, '4,5 h validadas')
        self.assertRedirects(self.client.get('/livro-de-campo/'), f'/livro-de-campo/{self.op.pk}/',
                             fetch_redirect_response=False)
        r = self.client.get(f'/livro-de-campo/{self.op.pk}/icnf.xlsx?ano=2026')
        self.assertEqual(r.status_code, 200)
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            wb = openpyxl.load_workbook(io.BytesIO(r.content))
        ws = wb['Acções de fogo Controlado']
        self.assertEqual((ws['B1'].value, ws['B2'].value), ('9-2020', 'Técnico tec'))
        self.assertEqual([ws[f'{c}5'].value for c in 'ABCDEF'],
                         ['L-12', '7', 'Alferce', 'Monchique', 'Faro', 'Queimada Autor.'])
        self.assertEqual((ws['G5'].value.date(), ws['H5'].value, ws['K5'].value, ws['M5'].value, ws['N5'].value,
                          ws['O5'].value, ws['P5'].value),
                         (date(2026, 1, 20), timedelta(hours=9), 'Mato', 'Responsável de queima', 2.5,
                          'Proprietário', 'J. Silva'))
        self.assertIsNone(ws['F6'].value)   # template placeholders cleared
        # Other users cannot export someone else's database
        other, _ = make_tecnico('outro', '10-2021')
        self.client.force_login(other)
        r = self.client.get(f'/livro-de-campo/{self.op.pk}/icnf.xlsx')
        self.assertEqual(r.status_code, 302)

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_validate_team(self):
        self.client.force_login(self.user)
        self.client.post(f'/registos/equipa/{self.ex.pk}/validar/')
        self.m_op.refresh_from_db()
        self.m_tec.refresh_from_db()
        self.assertEqual((self.m_op.validated_by, self.m_tec.validated_by), (self.user, None))
        plain = User.objects.create_user('plain', password='x')
        self.m_op.validated_by = None
        self.m_op.save()
        self.client.force_login(plain)
        self.client.post(f'/registos/equipa/{self.ex.pk}/validar/')
        self.m_op.refresh_from_db()
        self.assertIsNone(self.m_op.validated_by)


class DeadlineTests(Base):
    def test_icnf_windows(self):
        self.assertEqual(icnf_window_after(date(2026, 1, 20)), (date(2026, 2, 1), date(2026, 2, 15)))
        self.assertEqual(icnf_window_after(date(2026, 2, 10)), (date(2026, 4, 1), date(2026, 4, 15)))
        self.assertEqual(icnf_window_after(date(2026, 6, 30)), (date(2026, 12, 1), date(2026, 12, 15)))
        self.assertEqual(icnf_window_after(date(2026, 12, 3)), (date(2027, 2, 1), date(2027, 2, 15)))

    def kinds(self, today):
        return {(i['kind'], i['due'], i['status']) for i in all_deadlines(today)
                if i['kind'] != 'monitorizacao'}

    def test_poq_impact_incident_and_sent(self):
        self.assertEqual(self.kinds(date(2026, 2, 5)), {('poq', date(2026, 2, 15), 'urgente')})
        self.assertIn(('impacto', date(2027, 7, 31), 'proximo'), self.kinds(date(2027, 5, 1)))
        self.plan.incident, self.plan.incident_date = True, date(2026, 1, 20)
        self.plan.save()
        self.assertIn(('incidente', date(2026, 2, 4), 'atrasado'), self.kinds(date(2026, 2, 5)))
        self.plan.poq_sent_on = date(2026, 2, 1)
        self.plan.save()
        self.assertEqual(self.kinds(date(2026, 2, 5)), set())   # POQ-Impacto is beyond the 4-month horizon

    def test_credential_and_pfc_deadlines(self):
        _, op = make_tecnico('renova', '11-2022', valid_until=date(2026, 3, 31))
        items = [i for i in all_deadlines(date(2026, 2, 5)) if i['kind'] == 'credenciacao']
        self.assertEqual([(i['title'], i['due'], i['start']) for i in items],
                         [(op.name, date(2026, 3, 31), date(2026, 1, 29))])
        self.assertEqual([i for i in all_deadlines(date(2025, 6, 1)) if i['kind'] == 'credenciacao'], [])

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_page_and_mark_sent(self):
        self.client.force_login(self.user)
        with mock.patch('registos.deadlines.timezone.localdate', return_value=date(2026, 2, 5)):
            r = self.client.get('/prazos/')
        self.assertContains(r, 'Envio do POQ-Execução')
        self.assertEqual(r.context['scope'], 'meus')
        r = self.client.post(f'/registos/envio/{self.plan.pk}/', {'which': 'poq', 'date': '2026-02-03'})
        self.assertRedirects(r, f'/acoes/burning/{self.plan.pk}/', fetch_redirect_response=False)
        self.plan.refresh_from_db()
        self.assertEqual(self.plan.poq_sent_on, date(2026, 2, 3))
        future = (timezone.localdate() + timedelta(days=2)).isoformat()
        self.client.post(f'/registos/envio/{self.plan.pk}/', {'which': 'impacto', 'date': future})
        self.plan.refresh_from_db()
        self.assertIsNone(self.plan.impact_sent_on)
        self.client.post(f'/registos/envio/{self.plan.pk}/', {'which': 'poq', 'clear': '1'})
        self.plan.refresh_from_db()
        self.assertIsNone(self.plan.poq_sent_on)
        self.assertContains(self.client.get(f'/acoes/burning/{self.plan.pk}/'), 'Envio ao ICNF')

    def test_alerts_digest_once_per_stage(self):
        with mock.patch('registos.alerts.send_async_each') as send:
            self.assertEqual(send_deadline_alerts(date(2026, 2, 5), 'https://x.pt'), 1)
            subject, body, to = send.call_args[0]
            self.assertEqual(to, ['tec@example.com'])
            self.assertIn('POQ-Execução', body)
            self.assertIn('https://x.pt/acoes/burning/', body)
            self.assertEqual(send_deadline_alerts(date(2026, 2, 6)), 0)       # same stage: not repeated
            self.assertEqual(send_deadline_alerts(date(2026, 2, 13)), 1)      # last reminder (3 days)
            self.assertEqual(send_deadline_alerts(date(2026, 2, 16)), 1)      # overdue
        self.assertTrue(AlertLog.objects.filter(key__endswith=':atraso').exists())
        UserPreferences.objects.create(user=User.objects.create_user('quiet', email='q@example.com'),
                                       email_deadlines=False)
        self.action.responsible_user = User.objects.get(username='quiet')
        self.action.save()
        with mock.patch('registos.alerts.send_async_each') as send:
            self.assertEqual(send_deadline_alerts(date(2026, 7, 25)), 0)
            send.assert_not_called()
