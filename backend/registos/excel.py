"""Per-technician database file for ICNF (Regulamento, art. 7.º n.º 3), filled into the official
BD_FOGO-CONTROLADO template: sheet 1 "Acções de fogo Controlado", sheet 2 "Planos de Fogo Controlado"."""
import io
import warnings
from datetime import datetime, timedelta
from pathlib import Path

from fire_actions.models import FireAction
from pfc.models import PFC
from .fieldbook import ICNF_FORMATION, ICNF_FUNCTIONS, ICNF_VALIDATOR

TEMPLATE = Path(__file__).resolve().parent / 'data' / 'BD_FOGO-CONTROLADO.xlsx'
FIRST_ROW = 5
LAST_TEMPLATE_ROW = 998   # rows pre-filled with "…" placeholders in the template


def _td(t):
    return timedelta(hours=t.hour, minutes=t.minute) if t else None


# Dropdown lists, kept on the template's "Folha3" sheet. The template's own lists are outdated (AFN, CPE…)
# and its dropdowns use an Excel extension that openpyxl drops, so both are rewritten here from the app's values.
LISTS = [   # (Folha3 column, header, values, register columns that use it)
    ('A', 'Funções', list(dict.fromkeys(ICNF_FUNCTIONS.values())), 'M'),
    ('B', 'Responsável que valida', list(ICNF_VALIDATOR.values()), 'O'),
    ('C', 'Objetivos', [label for _, label in FireAction.OBJECTIVES]
     + [label for _, label in FireAction.QUEIMADA_PURPOSE], 'J'),
    ('D', 'Ação de fogo', ['Plano Fogo Controlado', 'Queimada Autor.', 'Formação'], 'F'),
    ('E', 'Formações vegetais', list(ICNF_FORMATION.values()), 'K'),
]


def _set_lists(wb, ws):
    from openpyxl.worksheet.datavalidation import DataValidation, DataValidationList
    lists = wb['Folha3']
    for row in lists.iter_rows(min_row=1, max_row=max(lists.max_row, 30), max_col=8):
        for cell in row:
            cell.value = None
    for col in ('IL', 'IN', 'IP', 'IR', 'IV'):   # the template's old lists, now unused
        for r in range(1, 11):
            ws[f'{col}{r}'].value = None
    ws.data_validations = DataValidationList()
    for col, header, values, target in LISTS:
        lists[f'{col}1'] = header
        for i, v in enumerate(values, start=2):
            lists[f'{col}{i}'] = v
        dv = DataValidation(type='list', formula1=f'Folha3!${col}$2:${col}${len(values) + 1}', allow_blank=True,
                            showErrorMessage=False)
        dv.add(f'{target}{FIRST_ROW}:{target}{LAST_TEMPLATE_ROW}')
        ws.add_data_validation(dv)


def generate_icnf_xlsx(operative, rows, today=None):
    import openpyxl
    with warnings.catch_warnings():   # the template has an Excel extension openpyxl does not read
        warnings.simplefilter('ignore')
        wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb['Acções de fogo Controlado']
    ws['B1'] = operative.credential_number
    ws['B2'] = operative.name
    _set_lists(wb, ws)

    # Clear the template placeholders ("…", 0, example note)
    for r in range(FIRST_ROW, LAST_TEMPLATE_ROW + 1):
        for col in 'ABCDEFGHIJKLMNOPQR':
            ws[f'{col}{r}'].value = None

    for i, e in enumerate(rows):
        r = FIRST_ROW + i
        values = {
            'A': e['book'], 'B': e['page'], 'C': e['freguesia'], 'D': e['concelho'], 'E': e['distrito'],
            'F': e['type'], 'G': datetime.combine(e['date'], datetime.min.time()), 'H': _td(e['start']),
            'I': _td(e['end']), 'J': e['objectives'], 'K': e['formation'], 'L': e['composition'],
            'M': e['function'], 'N': e['area'], 'O': e['validation_role'], 'P': e['validation_name'],
            'Q': e['other'], 'R': e['notes'],
        }
        for col, v in values.items():
            cell = ws[f'{col}{r}']
            cell.value = v if v not in ('', None) else None
            if col == 'G':
                cell.number_format = 'dd/mm/yy;@'
            elif col in 'HI':
                cell.number_format = '[h]:mm;@'

    # Sheet 2: the technician's PFCs
    ws2 = wb['Planos de Fogo Controlado']
    pfcs = PFC.objects.filter(technician=operative.user).prefetch_related('plan_parcels__parcel') \
        if operative.user_id else PFC.objects.none()
    for i, pfc in enumerate(pfcs.order_by('validity_start')):
        r = FIRST_ROW + i
        parcels = [pp.parcel for pp in pfc.plan_parcels.all()]
        executed = sum(p.area_ha or 0 for p in parcels if p.executions.filter(plan__pre_plan__pfc=pfc).exists())
        total = pfc.total_area_ha or 0
        approved = pfc.effective_approval_date
        for col, v in {
            'A': pfc.number, 'B': pfc.name,
            'C': ', '.join(sorted({p.freguesia for p in parcels if p.freguesia})),
            'D': ', '.join(sorted({p.concelho for p in parcels if p.concelho})),
            'E': ', '.join(sorted({p.distrito for p in parcels if p.distrito})),
            'F': pfc.proponent,
            'G': datetime.combine(approved, datetime.min.time()) if approved else None,
            'H': round(total, 2) if total else None,
            'I': round(executed / total, 4) if total else None,
        }.items():
            ws2[f'{col}{r}'] = v
        ws2[f'G{r}'].number_format = 'dd/mm/yy;@'
        ws2[f'I{r}'].number_format = '0%'

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
